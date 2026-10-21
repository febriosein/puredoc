import os
import sys
import io
import hashlib

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from PIL import Image, ImageDraw
import docx
from docx.shared import Inches
from fastapi.testclient import TestClient

from app.main import app
from app.config import STORAGE_DIR, TEMP_DIR
from app.docx_analyzer import analyze_docx, validate_docx_file
from app.word_converter import convert_docx_to_pdf_word
from app.pdf_restorer import restore_pdf_original_images
from app.verifier import verify_conversion

def create_sample_docx(target_path: str, img_dims=(1600, 1000)) -> str:
    # 1. Create high-res test image
    img = Image.new("RGB", img_dims, color=(24, 76, 120))
    draw = ImageDraw.Draw(img)
    draw.rectangle([50, 50, img_dims[0] - 50, img_dims[1] - 50], outline=(255, 215, 0), width=8)
    draw.text((100, 200), "PureDoc High-Fidelity Test Image", fill=(255, 255, 255))
    
    img_temp_path = os.path.join(TEMP_DIR, "temp_test_image.jpg")
    img.save(img_temp_path, "JPEG", quality=95)

    # 2. Build docx
    doc = docx.Document()
    doc.add_heading("PureDoc PDF Automated Verification", 0)
    p = doc.add_paragraph("Uji coba pipeline konversi dan verifikasi integritas resolusi asli.")
    doc.add_picture(img_temp_path, width=Inches(5.5))
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Kolom A"
    table.cell(0, 1).text = "Kolom B"
    table.cell(1, 0).text = "Nilai 1"
    table.cell(1, 1).text = "Nilai 2"

    os.makedirs(os.path.dirname(target_path), exist_ok=True)
    doc.save(target_path)
    return img_temp_path

def test_full_pipeline_end_to_end():
    sample_docx = os.path.join(TEMP_DIR, "test_document.docx")
    sample_pdf = os.path.join(TEMP_DIR, "test_document.pdf")
    
    orig_img_file = create_sample_docx(sample_docx, img_dims=(1600, 1000))
    
    with open(orig_img_file, "rb") as f:
        orig_bytes = f.read()
    orig_sha256 = hashlib.sha256(orig_bytes).hexdigest()

    # Step 1: Validate file
    assert validate_docx_file(sample_docx) is True

    # Step 2: Analyze docx
    analysis = analyze_docx(sample_docx)
    assert analysis["is_valid"] is True
    assert analysis["image_count"] >= 1
    assert analysis["min_resolution"] == "1600×1000 px"
    assert analysis["images"][0]["width"] == 1600
    assert analysis["images"][0]["height"] == 1000
    assert analysis["images"][0]["sha256"] == orig_sha256

    # Step 3: Native Word Conversion
    conv_result = convert_docx_to_pdf_word(sample_docx, sample_pdf)
    assert conv_result["success"] is True
    assert os.path.exists(sample_pdf)
    assert conv_result["word_page_count"] >= 1

    # Step 4: Lossless Image Restoration (Pass-through original raw bytes)
    restore_result = restore_pdf_original_images(sample_docx, sample_pdf)
    assert restore_result["total_images"] >= 1

    # Step 5: Post-conversion Verification
    verification = verify_conversion(sample_docx, sample_pdf, conv_result["word_page_count"], analysis)
    assert verification["status_code"] in ("LOLOS", "LOLOS_DENGAN_PERINGATAN")
    assert verification["metrics"]["word_pages"] == verification["metrics"]["pdf_pages"]
    assert len(verification["metrics"]["dimension_mismatches"]) == 0
    assert verification["quality_table"][0]["item"] == "Jumlah halaman"

    print("End-to-end pipeline test passed successfully!")

def test_fastapi_endpoints():
    client = TestClient(app)

    # 1. Health check
    res = client.get("/api/health")
    assert res.status_code == 200
    assert res.json()["status"] == "healthy"

    # 2. Upload docx
    sample_docx = os.path.join(TEMP_DIR, "test_api_doc.docx")
    create_sample_docx(sample_docx, img_dims=(1200, 800))
    
    with open(sample_docx, "rb") as f:
        res = client.post("/api/upload", files={"file": ("test_api_doc.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    job_id = data["job_id"]
    assert "analysis" in data

    # 3. Convert job
    res = client.post(f"/api/jobs/{job_id}/convert")
    assert res.status_code == 200
    conv_data = res.json()
    assert conv_data["success"] is True
    assert conv_data["verification"]["status_code"] in ("LOLOS", "LOLOS_DENGAN_PERINGATAN")

    # 4. Preview / Download PDF
    res = client.get(f"/api/jobs/{job_id}/pdf?download=false")
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"

    # 5. Get Quality Report
    res = client.get(f"/api/jobs/{job_id}/report")
    assert res.status_code == 200
    assert "verification" in res.json()

    # 6. Delete job (FR-20)
    res = client.delete(f"/api/jobs/{job_id}")
    assert res.status_code == 200
    assert res.json()["success"] is True

    print("All FastAPI endpoints tested successfully!")

if __name__ == "__main__":
    test_full_pipeline_end_to_end()
    test_fastapi_endpoints()
